"""Formatting helpers: sanitize, DOCX, PDF and HTML preview."""
import html
import io
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import FOOTER_TEXT, LOGO_PATH  # noqa: E402

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u2022": "-", "\u00a0": " ",
}
_HEADING_RE = re.compile(r"^(\d+\.\s+.+|[A-Z][A-Za-z &/,'-]{2,60}:|[A-Z0-9 &/,'-]{4,60})$")


def sanitize_text(text: str) -> str:
    """Remove typographic quotes / markdown noise so every exporter stays clean."""
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    text = text.replace("**", "")
    text = re.sub(r"^\s*#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*\*\s+", "- ", text, flags=re.MULTILINE)
    return text.strip()


def _is_heading(line: str) -> bool:
    line = line.strip()
    return bool(line) and len(line) < 70 and bool(_HEADING_RE.match(line)) and not line.endswith(".")


def _split_terms(terms: str):
    return [t.strip() for t in (terms or "").split(";") if t.strip()]


def _latin(s: str) -> str:
    return s.encode("latin-1", "ignore").decode("latin-1")


# ---------------------------------------------------------------- DOCX
def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    text = sanitize_text(text)
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    if os.path.exists(LOGO_PATH):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO_PATH, width=Inches(2.2))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.size = Pt(16)

    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        p = doc.add_paragraph()
        p.add_run(line).bold = _is_heading(line)
        p.paragraph_format.space_after = Pt(6)

    term_list = _split_terms(terms)
    if term_list:
        doc.add_paragraph().add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for i, t in enumerate(term_list, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = t

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = FOOTER_TEXT
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
class _LegalPDF(FPDF):
    def __init__(self, doc_type):
        super().__init__()
        self.doc_type = doc_type

    def header(self):
        if os.path.exists(LOGO_PATH):
            self.image(LOGO_PATH, x=(self.w - 45) / 2, y=8, w=45)
        self.set_y(30)
        self.set_font("Arial", "B", 13)
        self.cell(0, 8, _latin(self.doc_type), ln=True, align="C")
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font("Arial", "I", 8)
        self.cell(0, 8, f"{FOOTER_TEXT}   Page {self.page_no()}", align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    text = sanitize_text(text)
    pdf = _LegalPDF(doc_type)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    for line in text.split("\n"):
        line = line.strip()
        if not line:
            pdf.ln(3)
            continue
        pdf.set_font("Arial", "B" if _is_heading(line) else "", 11)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, 6, _latin(line))
        pdf.ln(1)

    out = pdf.output(dest="S")
    return out.encode("latin-1") if isinstance(out, str) else bytes(out)


# ---------------------------------------------------------------- HTML
def format_html_preview(text: str) -> str:
    text = sanitize_text(text)
    parts = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if _is_heading(line):
            parts.append(f"<h4 style='margin:16px 0 4px 0;color:#fff;'>{html.escape(line)}</h4>")
        elif line.startswith("- "):
            parts.append(f"<li style='margin-left:20px;'>{html.escape(line[2:])}</li>")
        else:
            parts.append(f"<p style='margin:4px 0;line-height:1.55;'>{html.escape(line)}</p>")
    return "".join(parts)
